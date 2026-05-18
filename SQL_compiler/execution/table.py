from typing import List, Dict, Any, Union
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import numpy as np


class Table:
    INTEGER = int
    FLOAT = float
    STRING = str
    DATE = datetime

    def __init__(self, name: str):
        self.name = name
        self.column_names: List[str] = []
        self.column_types: Dict[str, type] = {}
        self.rows: List[Dict[str, Any]] = []
        self.aliases: Dict[str, str] = {}

    def add_column(self, name: str, col_type: type) -> None:
        self.column_names.append(name)
        self.column_types[name] = col_type

    def add_row(self, row: Dict[str, Any]) -> None:
        for col in self.column_names:
            if col not in row:
                row[col] = None
        self.rows.append(row.copy())

    def get_value(self, row_idx: int, column: str) -> Any:
        if row_idx < 0 or row_idx >= len(self.rows):
            raise IndexError(f"Row index {row_idx} out of range")
        if column not in self.column_names:
            raise KeyError(f"Column '{column}' not found in table '{self.name}'")
        return self.rows[row_idx].get(column)

    def get_column_type(self, column: str) -> type:
        if column not in self.column_types:
            # Ищем без префикса
            base_name = column.split(".")[-1]
            if base_name in self.column_types:
                return self.column_types[base_name]
            raise KeyError(f"Column '{column}' not found in table '{self.name}'")
        return self.column_types[column]

    def get_column_index(self, column: str) -> int:
        if column not in self.column_names:
            raise KeyError(f"Column '{column}' not found in table '{self.name}'")
        return self.column_names.index(column)

    def __len__(self) -> int:
        return len(self.rows)

    def __str__(self) -> str:
        return f"Table('{self.name}', columns={len(self.column_names)}, rows={len(self.rows)})"

    def print_schema(self):
        print(f"\nТаблица: {self.name}")
        print("-" * 50)
        print(f"{'Колонка':<20} {'Тип':<15} {'Примеры':<30}")
        print("-" * 50)

        for col in self.column_names:
            col_type = self.column_types[col]
            type_name = {
                int: "INTEGER",
                float: "FLOAT",
                str: "STRING",
                datetime: "DATE",
            }.get(col_type, str(col_type))

            examples = []
            for i in range(min(3, len(self.rows))):
                val = self.rows[i].get(col)
                if val is not None:
                    if isinstance(val, datetime):
                        examples.append(val.strftime("%Y-%m-%d"))
                    else:
                        examples.append(str(val))

            example_str = ", ".join(examples) if examples else "(нет данных)"
            print(f"{col:<20} {type_name:<15} {example_str:<30}")


class ExcelLoader:
    NULL_VALUES = {None, "", "\\N", "NULL", "null", "None"}

    # Добавляем только это - единый формат для вывода
    DATE_OUTPUT_FORMAT = "%Y-%m-%d"

    def __init__(self):
        self.tables: Dict[str, Table] = {}

    def load(self, filepath: Union[str, Path]) -> Dict[str, Table]:
        filepath = Path(filepath)
        if not filepath.exists():
            raise FileNotFoundError(f"Excel file not found: {filepath}")

        print(f"Загрузка Excel файла: {filepath}")

        excel_file = pd.ExcelFile(filepath)

        for sheet_name in excel_file.sheet_names:
            print(f"  Обработка листа: {sheet_name}")
            df = pd.read_excel(filepath, sheet_name=sheet_name)
            table = self._create_table_from_dataframe(sheet_name, df)
            self.tables[sheet_name] = table
            table.print_schema()

        print(f"Загружено таблиц: {len(self.tables)}")
        return self.tables

    def _create_table_from_dataframe(self, name: str, df: pd.DataFrame) -> Table:
        table = Table(name)

        for col in df.columns:
            col_type = self._detect_column_type(df[col])
            table.add_column(col, col_type)
            print(f"    Колонка '{col}': {col_type.__name__}")

        for _, row in df.iterrows():
            row_dict = {}
            for col in df.columns:
                value = row[col]
                target_type = table.get_column_type(col)

                if pd.isna(value) or value == "" or value is None:
                    row_dict[col] = None
                else:
                    row_dict[col] = self._convert_value(value, target_type)
            table.add_row(row_dict)

        return table

    def _detect_column_type(self, series: pd.Series) -> type:
        non_null = series.dropna()
        non_null = non_null[non_null != ""]

        if len(non_null) == 0:
            return str

        # Проверяем на целые числа
        try:
            if all(isinstance(x, (int, np.integer)) or
                   (isinstance(x, float) and x == int(x)) or
                   (isinstance(x, str) and x.strip().lstrip('-').isdigit())
                   for x in non_null.head(100)):
                return int
        except:
            pass

        # Проверяем на числа с плавающей точкой
        try:
            if all(self._is_numeric(x) for x in non_null.head(100)):
                return float
        except:
            pass

        # Проверяем на даты
        if all(self._is_date(x) for x in non_null.head(100)):
            return datetime

        return str

    def _is_numeric(self, value) -> bool:
        if isinstance(value, (int, float, np.integer, np.floating)):
            return True
        if isinstance(value, str):
            try:
                float(value.strip())
                return True
            except:
                return False
        return False

    def _is_date(self, value) -> bool:
        if isinstance(value, (pd.Timestamp, datetime)):
            return True
        if isinstance(value, str):
            value = value.strip()
            for fmt in ["%Y-%m-%d", "%d.%m.%Y", "%Y/%m/%d", "%d-%m-%Y",
                        "%m/%d/%Y", "%d/%m/%Y", "%d.%m.%y", "%Y%m%d"]:
                try:
                    datetime.strptime(value, fmt)
                    return True
                except:
                    continue
        return False

    def _normalize_date(self, value: Any) -> Any:
        """Привести дату к единому формату YYYY-MM-DD"""
        if value is None:
            return None

        # Если уже в правильном формате
        if isinstance(value, str) and len(value) == 10 and value[4] == '-' and value[7] == '-':
            return value

        # Если datetime или Timestamp
        if isinstance(value, (pd.Timestamp, datetime)):
            dt = value.to_pydatetime() if isinstance(value, pd.Timestamp) else value
            return dt.strftime("%Y-%m-%d")

        # Если строка - парсим
        if isinstance(value, str):
            value = value.strip()
            for fmt in ["%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y",
                        "%Y/%m/%d", "%Y.%m.%d", "%m/%d/%Y", "%d.%m.%y", "%Y%m%d"]:
                try:
                    dt = datetime.strptime(value, fmt)
                    return dt.strftime("%Y-%m-%d")
                except ValueError:
                    continue
        return value

    def _convert_value(self, value: Any, target_type: type) -> Any:
        if value is None or (isinstance(value, str) and value.strip() == ""):
            return None

        try:
            if target_type == int:
                if isinstance(value, (float, np.floating)):
                    return int(value) if abs(value - int(value)) < 0.000001 else None
                if isinstance(value, str):
                    return int(float(value.strip()))
                return int(value)

            elif target_type == float:
                if isinstance(value, str):
                    return float(value.strip())
                return float(value)

            elif target_type == datetime:
                # Только здесь изменяем - нормализуем дату
                normalized = self._normalize_date(value)
                return normalized if normalized != value else None

            else:
                return str(value)
        except (ValueError, TypeError):
            return None
